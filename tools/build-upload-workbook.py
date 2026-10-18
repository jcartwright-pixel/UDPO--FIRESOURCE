"""Builds public/templates/Plant-Setup-Workbook.xlsx, the Upload Center's one workbook (Joe 10/18, his pick B), from the
Upload Center's own rules (functions/src/upload-rules.js SPEC), so the file and the checks never disagree.

    python3 tools/build-upload-workbook.py        (needs node and openpyxl)

Every tab: row 1 the column names (must-fill ones dark red with *), row 2 an example (grey; Lines has a second one on row 3;
the Upload Center skips them if left as is), drop-downs on choice columns on rows 2 to MAX_ROWS + 1 (the most rows the Upload
Center takes), and every column formatted as text (a column style, so Employee IDs keep leading zeros on every row). First tab: How to fill."""
import json, os, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SPEC = json.loads(subprocess.check_output(['node', '-e', 'const U=require("./functions/src/upload-rules.js");console.log(JSON.stringify({spec:U.SPEC,parts:U.PARTS,maxRows:U.MAX_ROWS}))'], cwd=ROOT))
NAVY, RED = '0B3D6B', '8B1A1A'
thin = Side(style='thin', color='C9D3DD')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
LAST = SPEC['maxRows'] + 1  # rows 2 to MAX_ROWS + 1: the most rows the Upload Center takes (upload-rules.js MAX_ROWS)
WRAP = Alignment(wrap_text=True, vertical='top')

wb = Workbook()
how = wb.active
how.title = 'How to fill'
how['A1'] = 'Operations - Plant Setup Workbook'
how['A1'].font = Font(size=16, bold=True, color=NAVY)
notes = [
    'One workbook for the whole plant. Each person fills their own tabs (the Who column below), then it is uploaded on Administration > Upload Center. Fill what you have; an empty tab is skipped and keeps what was saved before.',
    'Columns in dark red with * must be filled; a few more must be filled only sometimes (the Must fill? column says when). Row 2 on each tab is an example (Lines has two, rows 2 and 3): type over it or delete it. Pick the plant on the Upload Center screen; it is not typed in the file.',
    'Lists the app already has are drop-downs with the app\'s own words. Codes, login numbers and MOCREO / RedZone picks are set on their screens, never in this file. Location Areas feeds the Safety Walkthrough; the Safety Checklist\'s areas and sub-areas come from it.',
    'The Upload Center checks every row and shows what is new, what changed and any problems before anything is saved. Nothing is saved until you press Save.',
    'Trucks and trailers come from the United Dairy Fleet List and are assigned in the app, so there is no truck or trailer column. Never put Social Security numbers, birth dates or home addresses in this file.',
]
for i, t in enumerate(notes, 2):
    c = how.cell(i, 1, t); c.alignment = WRAP
    how.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
    how.row_dimensions[i].height = 36
row = len(notes) + 3
for j, t in enumerate(['Tab', 'Who fills it', 'Column', 'Must fill?', 'What to type', 'Where it goes in the app'], 1):
    c = how.cell(row, j, t); c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=NAVY); c.border = BORDER
for t in SPEC['spec']:
    ws = wb.create_sheet(t['tab'][:31])
    for i, c in enumerate(t['cols'], 1):
        cell = ws.cell(1, i, c['h'] + (' *' if c['req'] else ''))
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor=RED if c['req'] else NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center'); cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(34, len(c['h']) + 6))
        if c['choices']:
            # A column that takes several choices (comma between) only warns, so "Drinks, TEA" can be typed.
            dv = DataValidation(type='list', formula1='"' + ','.join(c['choices']) + '"', allow_blank=True, showErrorMessage=True,
                                errorStyle='warning' if c.get('multi') else 'stop',
                                errorTitle='Pick from the list', error=('Pick from: ' if c.get('multi') else 'Pick one of: ') + ', '.join(c['choices']))
            ws.add_data_validation(dv); dv.add(get_column_letter(i) + '2:' + get_column_letter(i) + str(LAST))
        row += 1
        what = c['what'] + (' Choices: ' + ', '.join(c['choices']) + '.' if c['choices'] else '') + (' Up to ' + str(c['max']) + ' letters.' if c.get('max') else '')
        must = 'Yes' if c['req'] else c.get('when', '')
        for j, v in enumerate([t['tab'], t['who'], c['h'], must, what.strip(), c['goes']], 1):
            cc = how.cell(row, j, v); cc.border = BORDER; cc.alignment = WRAP
            if j == 4 and must: cc.font = Font(bold=True, color=RED)
    # Text format on the whole column (a <col> style in the file, so it stays small); Excel and LibreOffice give that style to
    # every cell typed in the column, rows 2 to MAX_ROWS + 1 and past them, so Employee IDs keep their leading zeros.
    for i in range(1, len(t['cols']) + 1):
        ws.column_dimensions[get_column_letter(i)].number_format = '@'
    for n, ex in enumerate(t['examples'], 2):
        for i, v in enumerate(ex, 1):
            cell = ws.cell(n, i, v); cell.font = Font(italic=True, color='6A7885'); cell.number_format = '@'
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = 'A2'
    nc = len(t['cols']) + 2
    ws.cell(1, nc, t['about']).font = Font(italic=True, color='5A6B7D')
    ws.cell(1, nc).alignment = WRAP
    ws.column_dimensions[get_column_letter(nc)].width = 60
for col, w in zip('ABCDEF', (22, 16, 26, 11, 56, 44)):
    how.column_dimensions[col].width = w
how.freeze_panes = how.cell(len(notes) + 4, 1)
out = os.path.join(ROOT, 'public', 'templates', 'Plant-Setup-Workbook.xlsx')
wb.save(out)
print(out, len(SPEC['spec']), 'tabs')
